"""Deterministic editable PPTX generation from verified analysis modules."""

from collections.abc import Sequence
from io import BytesIO
from typing import Any, Protocol

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.presentation import Presentation as PresentationType
from pptx.util import Inches, Pt

from contract_models import AnalysisResult, PresentationRequest

FONT_FAMILY = "Microsoft JhengHei"
NAVY = RGBColor(15, 23, 42)
SLATE = RGBColor(71, 85, 105)
MUTED = RGBColor(226, 232, 240)
PANEL = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(20, 184, 166)
AMBER = RGBColor(217, 119, 6)


class PresentationGenerationError(RuntimeError):
    """Raised when verified module data cannot produce a valid PPTX."""


class PresentationGenerator(Protocol):
    """Application boundary for presentation generation."""

    def generate(
        self,
        request: PresentationRequest,
        modules: Sequence[AnalysisResult],
    ) -> bytes: ...


class PptxPresentationGenerator:
    """Build a small editable policy deck without another model call."""

    def generate(
        self,
        request: PresentationRequest,
        modules: Sequence[AnalysisResult],
    ) -> bytes:
        if not modules:
            raise PresentationGenerationError("At least one analysis module is required")

        try:
            deck = Presentation()
            deck.slide_width = Inches(13.333)
            deck.slide_height = Inches(7.5)
            self._add_cover(deck, request)
            for module in modules:
                self._add_module_slide(deck, module)
            self._add_trust_slide(deck, modules)

            output = BytesIO()
            deck.save(output)
            content = output.getvalue()
        except PresentationGenerationError:
            raise
        except (KeyError, TypeError, ValueError, OSError) as error:
            raise PresentationGenerationError(
                "Could not generate presentation"
            ) from error

        if not content.startswith(b"PK"):
            raise PresentationGenerationError("Generator returned an invalid PPTX")
        return content

    def _add_cover(
        self,
        deck: PresentationType,
        request: PresentationRequest,
    ) -> None:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        self._set_background(slide, NAVY)
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(0.18),
            Inches(7.5),
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = TEAL
        accent.line.fill.background()

        brand = slide.shapes.add_textbox(
            Inches(0.85), Inches(0.75), Inches(8.0), Inches(0.4)
        )
        self._set_text(
            brand.text_frame.paragraphs[0],
            "YOUTHLM  •  EVIDENCE WORKSPACE",
            13,
            bold=True,
            color=TEAL,
        )

        title = slide.shapes.add_textbox(
            Inches(0.85), Inches(1.65), Inches(11.5), Inches(1.45)
        )
        title.text_frame.word_wrap = True
        title.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        self._set_text(
            title.text_frame.paragraphs[0],
            request.title,
            34,
            bold=True,
            color=WHITE,
        )

        subtitle = slide.shapes.add_textbox(
            Inches(0.9), Inches(3.45), Inches(11.2), Inches(2.35)
        )
        lines = ["YouthLM｜可追溯政策分析簡報"]
        if request.audience:
            lines.append(f"對象：{request.audience}")
        lines.append(f"資料模組：{', '.join(request.source_module_ids)}")
        if request.instructions:
            lines.append(f"簡報要求：{request.instructions}")
        frame = subtitle.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            self._set_text(paragraph, line, 17, color=MUTED)
            paragraph.space_after = Pt(10)

        footer = slide.shapes.add_textbox(
            Inches(0.9), Inches(6.72), Inches(10.5), Inches(0.3)
        )
        self._set_text(
            footer.text_frame.paragraphs[0],
            "可編輯 PPTX  •  來源可追溯  •  分析限制不遺漏",
            11,
            color=MUTED,
        )

    def _add_module_slide(
        self,
        deck: PresentationType,
        module: AnalysisResult,
    ) -> None:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        self._set_background(slide, WHITE)
        self._add_title(slide, module.title)
        self._add_summary(slide, module.summary)

        if not self._add_chart(slide, module):
            self._add_table(slide, module)
        self._add_footer(slide, module.module_id)

    def _add_title(self, slide: Any, title: str) -> None:
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.55),
            Inches(0.24),
            Inches(0.12),
            Inches(0.58),
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = TEAL
        accent.line.fill.background()
        box = slide.shapes.add_textbox(
            Inches(0.82), Inches(0.22), Inches(11.8), Inches(0.65)
        )
        paragraph = box.text_frame.paragraphs[0]
        self._set_text(paragraph, title, 25, bold=True, color=NAVY)

    def _add_summary(self, slide: Any, summary: str) -> None:
        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6), Inches(1.1), Inches(3.4), Inches(5.55)
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = PANEL
        panel.line.color.rgb = MUTED
        frame = panel.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        heading = frame.paragraphs[0]
        self._set_text(heading, "分析摘要", 12, bold=True, color=TEAL)
        heading.space_after = Pt(12)
        paragraph = frame.add_paragraph()
        self._set_text(paragraph, summary, 16, color=NAVY)
        paragraph.line_spacing = 1.2

    def _add_chart(self, slide: Any, module: AnalysisResult) -> bool:
        spec = module.visualization
        if spec is None or spec.type not in {"line", "bar"}:
            return False
        if spec.x_field is None or spec.y_field is None:
            return False

        chart_data = self._build_chart_data(module)
        if chart_data is None:
            return False

        chart_type = (
            XL_CHART_TYPE.LINE_MARKERS
            if spec.type == "line"
            else XL_CHART_TYPE.COLUMN_CLUSTERED
        )
        chart = slide.shapes.add_chart(
            chart_type,
            Inches(4.15),
            Inches(1.15),
            Inches(8.6),
            Inches(5.65),
            chart_data,
        ).chart
        chart.style = 10
        chart.font.name = FONT_FAMILY
        chart.font.size = Pt(11)
        chart.has_title = True
        chart.chart_title.text_frame.text = (
            f"{spec.title}（{spec.unit}）" if spec.unit else spec.title
        )
        for paragraph in chart.chart_title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = FONT_FAMILY
                run.font.size = Pt(14)
                self._set_east_asian_font(run)
        chart.has_legend = len(chart_data) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        chart.value_axis.has_major_gridlines = True
        measure_values = [
            record.get(spec.y_field)
            for record in module.result_data.records
            if self._is_number(record.get(spec.y_field))
        ]
        if measure_values and min(measure_values) >= 0:
            chart.value_axis.minimum_scale = 0
        return True

    def _build_chart_data(self, module: AnalysisResult) -> ChartData | None:
        spec = module.visualization
        if spec is None or spec.x_field is None or spec.y_field is None:
            return None

        categories: list[Any] = []
        series_values: dict[tuple[Any, ...], dict[Any, int | float | None]] = {}
        series_fields = spec.series_fields or []

        for record in module.result_data.records:
            category = record.get(spec.x_field)
            value = record.get(spec.y_field)
            if category is None or not self._is_number_or_none(value):
                return None
            if category not in categories:
                categories.append(category)

            series_key = tuple(record.get(field) for field in series_fields)
            points = series_values.setdefault(series_key, {})
            if category in points:
                return None
            points[category] = value

        if not categories or not series_values:
            return None

        chart_data = ChartData()
        chart_data.categories = [str(category) for category in categories]
        y_label = self._column_label(module, spec.y_field)
        for series_key, points in series_values.items():
            label = " / ".join(str(value) for value in series_key) or y_label
            chart_data.add_series(label, [points.get(item) for item in categories])
        return chart_data

    def _add_table(self, slide: Any, module: AnalysisResult) -> None:
        columns = module.result_data.columns[:6]
        records = module.result_data.records[:10]
        if not columns or not records:
            box = slide.shapes.add_textbox(
                Inches(4.15), Inches(1.4), Inches(8.5), Inches(1.0)
            )
            self._set_text(
                box.text_frame.paragraphs[0],
                "此分析沒有可呈現的資料列。",
                16,
                color=SLATE,
            )
            return

        table = slide.shapes.add_table(
            len(records) + 1,
            len(columns),
            Inches(4.15),
            Inches(1.25),
            Inches(8.6),
            Inches(5.7),
        ).table
        for column_index, column in enumerate(columns):
            cell = table.cell(0, column_index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            paragraph = cell.text_frame.paragraphs[0]
            self._set_text(paragraph, column.label, 11, bold=True, color=WHITE)
            paragraph.alignment = PP_ALIGN.CENTER

        for row_index, record in enumerate(records, start=1):
            for column_index, column in enumerate(columns):
                value = record.get(column.name)
                cell = table.cell(row_index, column_index)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if row_index % 2 else PANEL
                paragraph = cell.text_frame.paragraphs[0]
                self._set_text(
                    paragraph,
                    "" if value is None else str(value),
                    10,
                    color=NAVY,
                )
                paragraph.alignment = PP_ALIGN.CENTER

    def _add_trust_slide(
        self,
        deck: PresentationType,
        modules: Sequence[AnalysisResult],
    ) -> None:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        self._set_background(slide, WHITE)
        self._add_title(slide, "資料來源與限制")
        box = slide.shapes.add_textbox(
            Inches(0.65), Inches(1.1), Inches(12.0), Inches(5.8)
        )
        frame = box.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.clear()

        lines: list[tuple[str, RGBColor]] = []
        for module in modules:
            for source in module.sources:
                agency = f"｜{source.agency}" if source.agency else ""
                lines.append(
                    (
                        f"來源｜{source.title}{agency}｜版本 {source.dataset_version_id}",
                        NAVY,
                    )
                )
            for warning in module.warnings:
                severity = {
                    "blocking": "阻擋",
                    "warning": "注意",
                    "info": "資訊",
                }[warning.severity]
                lines.append(
                    (
                        f"{severity}｜{warning.message}",
                        AMBER if warning.severity != "info" else SLATE,
                    )
                )
        if not lines:
            lines.append(("此簡報沒有來源或限制資訊。", SLATE))

        for index, (line, color) in enumerate(dict.fromkeys(lines)):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            self._set_text(paragraph, line, 15, color=color)
            paragraph.space_after = Pt(9)
        self._add_footer(slide, "Trust & Provenance")

    def _add_footer(self, slide: Any, label: str) -> None:
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.6),
            Inches(7.05),
            Inches(12.1),
            Inches(0.015),
        )
        line.fill.solid()
        line.fill.fore_color.rgb = MUTED
        line.line.fill.background()
        box = slide.shapes.add_textbox(
            Inches(0.65), Inches(7.1), Inches(12.0), Inches(0.2)
        )
        self._set_text(
            box.text_frame.paragraphs[0],
            f"YouthLM  •  {label}",
            8,
            color=SLATE,
        )

    @staticmethod
    def _set_background(slide: Any, color: RGBColor) -> None:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color

    @staticmethod
    def _set_text(
        paragraph: Any,
        text: str,
        size: int,
        *,
        bold: bool = False,
        color: RGBColor = NAVY,
    ) -> None:
        paragraph.text = text
        for run in paragraph.runs:
            run.font.name = FONT_FAMILY
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            PptxPresentationGenerator._set_east_asian_font(run)

    @staticmethod
    def _set_east_asian_font(run: Any) -> None:
        run_properties = run._r.get_or_add_rPr()
        east_asian = run_properties.find(qn("a:ea"))
        if east_asian is None:
            east_asian = OxmlElement("a:ea")
            run_properties.append(east_asian)
        east_asian.set("typeface", FONT_FAMILY)

    @staticmethod
    def _column_label(module: AnalysisResult, name: str) -> str:
        for column in module.result_data.columns:
            if column.name == name:
                return column.label
        return name

    @staticmethod
    def _is_number_or_none(value: Any) -> bool:
        return value is None or PptxPresentationGenerator._is_number(value)

    @staticmethod
    def _is_number(value: Any) -> bool:
        return (
            isinstance(value, (int, float)) and not isinstance(value, bool)
        )
