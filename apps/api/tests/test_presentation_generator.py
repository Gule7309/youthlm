"""Tests for deterministic editable PPTX generation."""

import json
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation

from contract_models import AnalysisResult, PresentationRequest
from presentation_generator import PptxPresentationGenerator

REPOSITORY_ROOT = Path(__file__).parents[3]
RESULT_FIXTURE = (
    REPOSITORY_ROOT
    / "contracts/fixtures/frontend-integration/analysis-result.example.json"
)


def analysis_result() -> AnalysisResult:
    return AnalysisResult.model_validate_json(
        RESULT_FIXTURE.read_text(encoding="utf-8")
    )


def presentation_request(result: AnalysisResult) -> PresentationRequest:
    return PresentationRequest(
        contract_version="0.1.0",
        project_id=result.project_id,
        source_module_ids=[result.module_id],
        title="板橋區青年人口趨勢",
        audience="青年政策規劃人員",
        output_format="pptx",
        instructions="保留資料限制。",
    )


def slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.text
    )


class PptxPresentationGeneratorTests(unittest.TestCase):
    def test_builds_editable_chart_and_trust_slides(self) -> None:
        result = analysis_result()

        content = PptxPresentationGenerator().generate(
            presentation_request(result),
            [result],
        )

        self.assertTrue(content.startswith(b"PK"))
        with ZipFile(BytesIO(content)) as archive:
            cover_xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
        self.assertIn("板橋區青年人口趨勢", cover_xml)
        self.assertIn('a:ea typeface="Microsoft JhengHei"', cover_xml)
        deck = Presentation(BytesIO(content))
        self.assertEqual(len(deck.slides), 3)
        self.assertIn("板橋區青年人口趨勢", slide_text(deck.slides[0]))
        self.assertIn("青年政策規劃人員", slide_text(deck.slides[0]))
        chart_shapes = [shape for shape in deck.slides[1].shapes if shape.has_chart]
        self.assertEqual(len(chart_shapes), 1)
        self.assertEqual(chart_shapes[0].chart.value_axis.minimum_scale, 0)
        trust_text = slide_text(deck.slides[2])
        self.assertIn("現住人口之年齡分配", trust_text)
        self.assertIn(result.warnings[0].message, trust_text)

    def test_falls_back_to_an_editable_table_without_chart_spec(self) -> None:
        result = analysis_result().model_copy(update={"visualization": None})

        content = PptxPresentationGenerator().generate(
            presentation_request(result),
            [result],
        )

        deck = Presentation(BytesIO(content))
        tables = [shape.table for shape in deck.slides[1].shapes if shape.has_table]
        self.assertEqual(len(tables), 1)
        table = tables[0]
        self.assertEqual(table.cell(0, 0).text, "年份")
        self.assertEqual(table.cell(1, 0).text, "2022")

    def test_generated_deck_contains_only_supplied_module_evidence(self) -> None:
        result = analysis_result()
        payload = json.loads(RESULT_FIXTURE.read_text(encoding="utf-8"))
        payload["module_id"] = "analysis_second"
        payload["title"] = "第二個分析模組"
        second = AnalysisResult.model_validate(payload)
        request = presentation_request(result).model_copy(
            update={"source_module_ids": [result.module_id, second.module_id]}
        )

        content = PptxPresentationGenerator().generate(request, [result, second])

        deck = Presentation(BytesIO(content))
        self.assertEqual(len(deck.slides), 4)
        self.assertIn(result.title, slide_text(deck.slides[1]))
        self.assertIn(second.title, slide_text(deck.slides[2]))


if __name__ == "__main__":
    unittest.main()
