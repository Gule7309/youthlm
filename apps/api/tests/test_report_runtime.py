"""HTTP integration tests for Report Artifact generation and download."""

import asyncio
import tempfile
import unittest
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import httpx
from docx import Document

from contract_models import AnalysisResult
from main import create_app
from module_store import InMemoryModuleStore, SQLiteModuleStore
from report_generator import DocxReportGenerator, ReportGenerationError
from report_result_store import SQLiteReportResultStore
from report_service import DOCX_MEDIA_TYPE, ReportService
from report_store import LocalReportArtifactStore, ReportArtifactStoreError

REPOSITORY_ROOT = Path(__file__).parents[3]
RESULT_FIXTURE = (
    REPOSITORY_ROOT
    / "contracts/fixtures/frontend-integration/analysis-result.example.json"
)


def request(app, method: str, path: str, *, json: dict | None = None) -> httpx.Response:
    async def send() -> httpx.Response:
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(
            transport=transport,
            base_url="http://test",
        ) as client:
            return await client.request(method, path, json=json)

    return asyncio.run(send())


def stored_result() -> AnalysisResult:
    return AnalysisResult.model_validate_json(
        RESULT_FIXTURE.read_text(encoding="utf-8")
    )


def report_request(result: AnalysisResult, **overrides) -> dict:
    payload = {
        "contract_version": "0.1.0",
        "project_id": result.project_id,
        "source_module_ids": [result.module_id],
        "title": "板橋區青年人口議題研析報告",
        "audience": "青年政策規劃人員",
        "language": "zh-TW",
        "template_id": "youthlm_default",
        "output_format": "docx",
        "instructions": "保留所有資料限制與來源。",
    }
    payload.update(overrides)
    return payload


class FailingGenerator:
    def generate(self, request, modules) -> bytes:
        raise ReportGenerationError(
            f"private generator detail: {request.title}/{len(modules)}"
        )


class FailingArtifactStore:
    def save(self, project_id: str, report_id: str, content: bytes):
        raise ReportArtifactStoreError(
            f"private storage detail: {project_id}/{report_id}/{len(content)}"
        )

    def get(self, project_id: str, report_id: str):
        raise ReportArtifactStoreError(
            f"private storage detail: {project_id}/{report_id}"
        )


class ReportRuntimeTests(unittest.TestCase):
    def test_generates_contract_result_then_downloads_real_docx(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result = stored_result()
            database_path = Path(directory) / "youthlm.sqlite3"
            artifact_directory = Path(directory) / "artifacts"
            module_store = SQLiteModuleStore(database_path)
            module_store.save(result)
            service = ReportService(
                module_store,
                DocxReportGenerator(),
                LocalReportArtifactStore(artifact_directory),
                result_store=SQLiteReportResultStore(database_path),
                id_factory=lambda: "report_test",
                clock=lambda: datetime(2026, 9, 10, 8, 0, tzinfo=UTC),
            )
            app = create_app(module_store=module_store, report_service=service)

            response = request(
                app,
                "POST",
                "/v1/reports",
                json=report_request(result),
            )

            self.assertEqual(response.status_code, 201)
            payload = response.json()
            self.assertEqual(payload["status"], "ready")
            self.assertEqual(payload["report_id"], "report_test")
            self.assertEqual(payload["output_format"], "docx")
            self.assertGreater(payload["file_size_bytes"], 1_000)
            self.assertEqual(len(payload["artifact_sha256"]), 64)
            self.assertEqual(
                payload["warnings"],
                [
                    warning.model_dump(mode="json", exclude_none=True)
                    for warning in result.warnings
                ],
            )

            restarted_store = SQLiteModuleStore(database_path)
            restarted_service = ReportService(
                restarted_store,
                DocxReportGenerator(),
                LocalReportArtifactStore(artifact_directory),
                result_store=SQLiteReportResultStore(database_path),
            )
            restarted_app = create_app(
                module_store=restarted_store,
                report_service=restarted_service,
            )
            download = request(restarted_app, "GET", payload["download_url"])
            self.assertEqual(download.status_code, 200)
            self.assertEqual(download.headers["content-type"], DOCX_MEDIA_TYPE)
            self.assertIn("report_test.docx", download.headers["content-disposition"])
            self.assertEqual(len(download.content), payload["file_size_bytes"])
            self.assertTrue(download.content.startswith(b"PK"))

            document = Document(BytesIO(download.content))
            paragraph_text = "\n".join(item.text for item in document.paragraphs)
            self.assertIn("板橋區青年人口議題研析報告", paragraph_text)
            self.assertIn(result.summary, paragraph_text)
            self.assertIn("資料來源與方法", paragraph_text)
            self.assertGreaterEqual(len(document.tables), 1)
            self.assertEqual(
                restarted_service.get_result(
                    result.project_id,
                    "report_test",
                ).model_dump(mode="json", exclude_none=True),
                payload,
            )

    def test_does_not_read_same_module_id_from_another_project(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result = stored_result()
            module_store = InMemoryModuleStore()
            module_store.save(result)
            service = ReportService(
                module_store,
                DocxReportGenerator(),
                LocalReportArtifactStore(directory),
            )
            response = request(
                create_app(module_store=module_store, report_service=service),
                "POST",
                "/v1/reports",
                json=report_request(result, project_id="different_project"),
            )

            self.assertEqual(response.status_code, 404)
            self.assertEqual(response.json()["error"]["code"], "module_not_found")
            self.assertEqual(
                response.json()["error"]["details"]["missing_module_ids"],
                [result.module_id],
            )

    def test_rejects_blocked_module_as_report_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result = stored_result().model_copy(update={"status": "blocked"})
            module_store = InMemoryModuleStore()
            module_store.save(result)
            service = ReportService(
                module_store,
                DocxReportGenerator(),
                LocalReportArtifactStore(directory),
            )
            response = request(
                create_app(module_store=module_store, report_service=service),
                "POST",
                "/v1/reports",
                json=report_request(result),
            )

            self.assertEqual(response.status_code, 422)
            self.assertEqual(response.json()["error"]["code"], "dataset_error")
            self.assertEqual(
                response.json()["error"]["details"]["blocked_module_ids"],
                [result.module_id],
            )

    def test_returns_safe_errors_for_generator_and_storage_failures(self) -> None:
        result = stored_result()
        module_store = InMemoryModuleStore()
        module_store.save(result)
        with tempfile.TemporaryDirectory() as directory:
            generator_service = ReportService(
                module_store,
                FailingGenerator(),
                LocalReportArtifactStore(directory),
            )
            generator_response = request(
                create_app(module_store=module_store, report_service=generator_service),
                "POST",
                "/v1/reports",
                json=report_request(result),
            )
            storage_service = ReportService(
                module_store,
                DocxReportGenerator(),
                FailingArtifactStore(),
            )
            storage_app = create_app(
                module_store=module_store,
                report_service=storage_service,
            )
            storage_response = request(
                storage_app,
                "POST",
                "/v1/reports",
                json=report_request(result),
            )
            download_response = request(
                storage_app,
                "GET",
                "/v1/projects/project_1/reports/report_1/download",
            )

        self.assertEqual(generator_response.status_code, 500)
        self.assertEqual(
            generator_response.json()["error"]["message"],
            "Report generation failed",
        )
        for response in (storage_response, download_response):
            self.assertEqual(response.status_code, 500)
            self.assertEqual(
                response.json()["error"]["message"],
                "Report artifact storage failed",
            )

    def test_invalid_request_uses_report_validation_message(self) -> None:
        response = request(
            create_app(module_store=InMemoryModuleStore()),
            "POST",
            "/v1/reports",
            json={
                "contract_version": "0.1.0",
                "project_id": "project_1",
                "source_module_ids": [],
                "title": "",
                "output_format": "docx",
            },
        )

        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["error"]["code"], "validation_error")
        self.assertEqual(
            response.json()["error"]["message"],
            "Report request validation failed",
        )

    def test_missing_or_cross_project_artifact_is_not_downloadable(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            service = ReportService(
                InMemoryModuleStore(),
                DocxReportGenerator(),
                LocalReportArtifactStore(directory),
            )
            response = request(
                create_app(report_service=service),
                "GET",
                "/v1/projects/project_other/reports/report_missing/download",
            )

            self.assertEqual(response.status_code, 404)
            self.assertEqual(response.json()["error"]["code"], "module_not_found")


if __name__ == "__main__":
    unittest.main()
