"""Deterministic editable DOCX generation from verified analysis modules."""

import json
from collections.abc import Sequence
from io import BytesIO
from typing import Protocol

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from contract_models import AnalysisResult, ReportRequest

FONT_FAMILY = "Microsoft JhengHei"
NAVY = RGBColor(15, 23, 42)
SLATE = RGBColor(71, 85, 105)
TEAL = RGBColor(13, 148, 136)
WHITE_HEX = "FFFFFF"
NAVY_HEX = "0F172A"
PANEL_HEX = "F1F5F9"
MAX_TABLE_COLUMNS = 6
MAX_TABLE_ROWS = 20


class ReportGenerationError(RuntimeError):
    """Raised when verified module data cannot produce a valid DOCX."""


class ReportGenerator(Protocol):
    """Application boundary for report generation."""

    def generate(
        self,
        request: ReportRequest,
        modules: Sequence[AnalysisResult],
    ) -> bytes: ...


class DocxReportGenerator:
    """Build an editable, traceable policy research report without a model call."""

    def generate(
        self,
        request: ReportRequest,
        modules: Sequence[AnalysisResult],
    ) -> bytes:
        if not modules:
            raise ReportGenerationError("At least one analysis module is required")
        try:
            document = Document()
            self._configure_document(document)
            self._add_cover(document, request)
            self._add_executive_summary(document, modules)
            for module in modules:
                self._add_module(document, module)
            self._add_methodology(document, modules)

            output = BytesIO()
            document.save(output)
            content = output.getvalue()
        except ReportGenerationError:
            raise
        except (KeyError, TypeError, ValueError, OSError) as error:
            raise ReportGenerationError("Could not generate report") from error
        if not content.startswith(b"PK"):
            raise ReportGenerationError("Generator returned an invalid DOCX")
        return content

    def _configure_document(self, document: DocumentType) -> None:
        section = document.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        normal = document.styles["Normal"]
        normal.font.name = FONT_FAMILY
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = NAVY
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FAMILY)

    def _add_cover(self, document: DocumentType, request: ReportRequest) -> None:
        brand = document.add_paragraph()
        brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._add_run(brand, "YOUTHLM  •  EVIDENCE WORKSPACE", 11, TEAL, bold=True)
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_before = Pt(72)
        self._add_run(title, request.title, 26, NAVY, bold=True)
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_before = Pt(18)
        self._add_run(subtitle, "議題研析報告｜可編輯 DOCX", 14, SLATE)
        details = [f"資料模組：{', '.join(request.source_module_ids)}"]
        if request.audience:
            details.append(f"適用對象：{request.audience}")
        if request.instructions:
            details.append(f"編製要求：{request.instructions}")
        for detail in details:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._add_run(paragraph, detail, 10, SLATE)
        note = document.add_paragraph()
        note.paragraph_format.space_before = Pt(72)
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._add_run(
            note,
            "本報告由已驗證的結構化分析成果產生；來源、版本與限制均保留。",
            9,
            SLATE,
        )
        document.add_page_break()

    def _add_executive_summary(
        self,
        document: DocumentType,
        modules: Sequence[AnalysisResult],
    ) -> None:
        self._add_heading(document, "執行摘要", level=1)
        for module in modules:
            paragraph = document.add_paragraph(style="List Bullet")
            self._add_run(paragraph, f"{module.title}：", 10.5, NAVY, bold=True)
            self._add_run(paragraph, module.summary, 10.5, NAVY)
        warning_count = sum(len(module.warnings) for module in modules)
        paragraph = document.add_paragraph()
        self._add_run(
            paragraph,
            f"本報告共引用 {len(modules)} 個分析模組，包含 {warning_count} 項資料限制或提醒。",
            9.5,
            SLATE,
        )

    def _add_module(self, document: DocumentType, module: AnalysisResult) -> None:
        self._add_heading(document, module.title, level=1)
        self._add_label_value(document, "研究問題", module.question)
        self._add_label_value(document, "分析狀態", module.status)
        self._add_label_value(
            document,
            "資料範圍",
            json.dumps(module.filters, ensure_ascii=False, sort_keys=True),
        )

        self._add_heading(document, "主要發現", level=2)
        document.add_paragraph(module.summary)
        self._add_heading(document, "結構化資料", level=2)
        self._add_result_table(document, module)

        self._add_heading(document, "資料限制", level=2)
        if module.warnings:
            for warning in module.warnings:
                paragraph = document.add_paragraph(style="List Bullet")
                severity = {
                    "blocking": "阻擋",
                    "warning": "注意",
                    "info": "資訊",
                }[warning.severity]
                self._add_run(paragraph, f"[{severity}] ", 10, NAVY, bold=True)
                self._add_run(paragraph, warning.message, 10, NAVY)
        else:
            document.add_paragraph("此分析模組未回報額外限制。")

    def _add_result_table(
        self,
        document: DocumentType,
        module: AnalysisResult,
    ) -> None:
        columns = module.result_data.columns[:MAX_TABLE_COLUMNS]
        records = module.result_data.records[:MAX_TABLE_ROWS]
        if not columns or not records:
            document.add_paragraph("此分析沒有可呈現的資料列。")
            return
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        table.autofit = True
        for index, column in enumerate(columns):
            cell = table.rows[0].cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            self._shade_cell(cell, NAVY_HEX)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label = f"{column.label}（{column.unit}）" if column.unit else column.label
            self._add_run(paragraph, label, 9, RGBColor(255, 255, 255), bold=True)
        for row_index, record in enumerate(records, start=1):
            cells = table.add_row().cells
            for column_index, column in enumerate(columns):
                cell = cells[column_index]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index % 2 == 0:
                    self._shade_cell(cell, PANEL_HEX)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                value = record.get(column.name)
                self._add_run(
                    paragraph,
                    "" if value is None else str(value),
                    8.5,
                    NAVY,
                )
        if len(module.result_data.records) > MAX_TABLE_ROWS:
            document.add_paragraph(
                f"表格僅顯示前 {MAX_TABLE_ROWS} 筆；完整資料仍保留於原分析模組。"
            )

    def _add_methodology(
        self,
        document: DocumentType,
        modules: Sequence[AnalysisResult],
    ) -> None:
        document.add_page_break()
        self._add_heading(document, "資料來源與方法", level=1)
        seen_sources: set[tuple[str, str]] = set()
        for module in modules:
            for source in module.sources:
                key = (source.source_id, source.dataset_version_id)
                if key in seen_sources:
                    continue
                seen_sources.add(key)
                parts = [source.title]
                if source.agency:
                    parts.append(source.agency)
                parts.append(f"版本 {source.dataset_version_id}")
                if source.source_url:
                    parts.append(source.source_url)
                document.add_paragraph("｜".join(parts), style="List Bullet")
        self._add_heading(document, "查詢與可追溯紀錄", level=2)
        for module in modules:
            for record in module.provenance:
                parameters = json.dumps(
                    record.query_parameters,
                    ensure_ascii=False,
                    sort_keys=True,
                )
                document.add_paragraph(
                    f"{module.title}｜{record.query_tool}｜{parameters}",
                    style="List Bullet",
                )
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(18)
        self._add_run(
            paragraph,
            "生成原則：數值直接來自 AnalysisResult.result_data；本文未由模型重新計算或補值。",
            9,
            SLATE,
        )

    def _add_heading(
        self,
        document: DocumentType,
        text: str,
        *,
        level: int,
    ) -> None:
        paragraph = document.add_heading(level=level)
        self._add_run(
            paragraph,
            text,
            18 if level == 1 else 13,
            NAVY if level == 1 else TEAL,
            bold=True,
        )

    def _add_label_value(
        self,
        document: DocumentType,
        label: str,
        value: str,
    ) -> None:
        paragraph = document.add_paragraph()
        self._add_run(paragraph, f"{label}：", 10, SLATE, bold=True)
        self._add_run(paragraph, value, 10, NAVY)

    @staticmethod
    def _shade_cell(cell: object, fill: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    @staticmethod
    def _add_run(
        paragraph: object,
        text: str,
        size: float,
        color: RGBColor,
        *,
        bold: bool = False,
    ) -> None:
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.name = FONT_FAMILY
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FAMILY)
